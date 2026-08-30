"""Deterministic CSV/XLSX budget parsers. Workbook bytes and cells are never logged."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.investment_planning.drive_binding import (
    CSV_MIME,
    MAX_BUDGET_BYTES,
    MAX_BUDGET_ROWS,
    SUPPORTED_BUDGET_MIMES,
    XLSX_MIME,
)
from app.investment_planning.errors import BudgetFormatError
from app.investment_planning.privacy import redact_customer_amounts


@dataclass(frozen=True, slots=True)
class ParsedAmount:
    raw: str
    value: Decimal | None
    missing: bool


@dataclass(frozen=True, slots=True)
class ParsedBudgetRow:
    row_index: int
    cells: dict[str, str]


@dataclass(frozen=True, slots=True)
class ParsedBudgetTable:
    headers: tuple[str, ...]
    rows: tuple[ParsedBudgetRow, ...]
    mime_type: str


def parse_budget_bytes(*, data: bytes, mime_type: str, file_name: str = "") -> ParsedBudgetTable:
    if len(data) > MAX_BUDGET_BYTES:
        raise BudgetFormatError("Budget file exceeds the maximum allowed size.")
    lowered = mime_type.lower()
    name = file_name.lower()
    if lowered not in SUPPORTED_BUDGET_MIMES and not name.endswith((".csv", ".xlsx")):
        raise BudgetFormatError("Only CSV and XLSX budget files are supported.")
    if lowered == CSV_MIME or name.endswith(".csv"):
        return _parse_csv(data)
    if lowered == XLSX_MIME or name.endswith(".xlsx"):
        return _parse_xlsx(data)
    raise BudgetFormatError("Only CSV and XLSX budget files are supported.")


def parse_decimal_cell(raw: str) -> ParsedAmount:
    text = raw.strip()
    if text == "":
        return ParsedAmount(raw=raw, value=None, missing=True)
    normalized = text.replace(",", "")
    try:
        value = Decimal(normalized)
    except InvalidOperation as exc:
        raise BudgetFormatError("Quarter cells must be numeric, blank, or zero.") from exc
    if value.is_nan() or value.is_infinite():
        raise BudgetFormatError("Quarter cells must be finite decimals.")
    return ParsedAmount(raw=raw, value=value, missing=False)


def safe_parse_error(exc: Exception) -> str:
    return redact_customer_amounts(type(exc).__name__)


def _parse_csv(data: bytes) -> ParsedBudgetTable:
    text = data.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = tuple(cell.strip() for cell in next(reader))
    except StopIteration as exc:
        raise BudgetFormatError("Budget CSV is missing a header row.") from exc
    if not headers or not any(headers):
        raise BudgetFormatError("Budget CSV is missing a header row.")
    rows: list[ParsedBudgetRow] = []
    for index, raw_row in enumerate(reader, start=2):
        if index - 1 > MAX_BUDGET_ROWS:
            raise BudgetFormatError("Budget file exceeds the maximum allowed rows.")
        if not any(cell.strip() for cell in raw_row):
            continue
        cells = {
            headers[col]: raw_row[col] if col < len(raw_row) else ""
            for col in range(len(headers))
            if headers[col]
        }
        rows.append(ParsedBudgetRow(row_index=index, cells=cells))
    return ParsedBudgetTable(headers=headers, rows=tuple(rows), mime_type=CSV_MIME)


def _parse_xlsx(data: bytes) -> ParsedBudgetTable:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = workbook.worksheets
        if len(sheets) != 1:
            raise BudgetFormatError("XLSX budget files must contain exactly one data sheet.")
        sheet = sheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            raise BudgetFormatError("Budget XLSX is missing a header row.")
        headers = tuple(str(cell).strip() if cell is not None else "" for cell in header_row)
        if not any(headers):
            raise BudgetFormatError("Budget XLSX is missing a header row.")
        rows: list[ParsedBudgetRow] = []
        for index, raw_row in enumerate(iterator, start=2):
            if index - 1 > MAX_BUDGET_ROWS:
                raise BudgetFormatError("Budget file exceeds the maximum allowed rows.")
            values = tuple(raw_row or ())
            if not any(value is not None and str(value).strip() for value in values):
                continue
            cells: dict[str, str] = {}
            for col, header in enumerate(headers):
                if not header:
                    continue
                cell: Any = values[col] if col < len(values) else None
                cells[header] = "" if cell is None else str(cell)
            rows.append(ParsedBudgetRow(row_index=index, cells=cells))
        return ParsedBudgetTable(headers=headers, rows=tuple(rows), mime_type=XLSX_MIME)
    finally:
        workbook.close()
